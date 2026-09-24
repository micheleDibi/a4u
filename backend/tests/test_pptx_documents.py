"""PPTX fra i documenti del corso (U3, WP2f).

Upload accettato solo se il contenuto è davvero una presentazione OOXML
(il MIME arriva dal client), testo delle slide per il riassunto, figure
estratte slide per slide con la riga «Fonte» che cita la slide.
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path
from typing import Any

import pytest
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.core.errors import ValidationAppError
from app.models.course_document_figure import CourseDocumentFigure
from app.services import course_document_figures_worker as worker
from app.services import document_extraction_service, file_service, remote_storage
from app.services.document_figures import office
from app.services.figure_attribution import figure_attribution_line
from tests.course_builders import build_course, build_course_document
from tests.fixtures.source_figures.build import GROUND_TRUTH, PPTX_NAME, build_all
from tests.test_source_figure_worker import FakeStorage, FakeVision

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


@pytest.fixture(scope="module")
def pptx_bytes(tmp_path_factory: pytest.TempPathFactory) -> bytes:
    directory = tmp_path_factory.mktemp("pptx")
    build_all(directory)
    return (directory / PPTX_NAME).read_bytes()


@pytest.fixture
def storage(monkeypatch: pytest.MonkeyPatch) -> FakeStorage:
    fake = FakeStorage()
    monkeypatch.setattr(remote_storage, "get_storage", lambda: fake)
    return fake


async def test_valid_pptx_upload_is_accepted(storage: FakeStorage, pptx_bytes: bytes) -> None:
    path, name, size = await file_service.save_document_from_bytes(
        pptx_bytes, subdir="courses/abc", mime_type=PPTX_MIME
    )
    assert name.endswith(".pptx") and size == len(pptx_bytes)
    assert remote_storage.uploads_key(path) in storage.files


@pytest.mark.parametrize(
    "payload",
    [b"%PDF-1.4 non una presentazione", b"PK\x03\x04 zip rotto", b"testo qualsiasi"],
)
async def test_fake_pptx_upload_is_refused(storage: FakeStorage, payload: bytes) -> None:
    with pytest.raises(ValidationAppError) as info:
        await file_service.save_document_from_bytes(
            payload, subdir="courses/abc", mime_type=PPTX_MIME
        )
    assert info.value.code == "invalid_document_content"
    assert storage.files == {}


async def test_docx_zip_declared_as_pptx_is_refused(storage: FakeStorage, tmp_path: Path) -> None:
    build_all(tmp_path)
    docx = (tmp_path / "vibrometria_appunti.docx").read_bytes()
    with pytest.raises(ValidationAppError):
        await file_service.save_document_from_bytes(docx, subdir="courses/abc", mime_type=PPTX_MIME)


def _rewrite(pptx: bytes, changes: dict[str, bytes]) -> bytes:
    """Copia del pacchetto con alcune parti sostituite o aggiunte."""
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(pptx)) as src,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst,
    ):
        for info in src.infolist():
            if info.filename not in changes:
                dst.writestr(info.filename, src.read(info))
        for name, data in changes.items():
            dst.writestr(name, data)
    return out.getvalue()


def _part(pptx: bytes, name: str) -> str:
    with zipfile.ZipFile(io.BytesIO(pptx)) as zf:
        return zf.read(name).decode("utf-8")


async def test_zip_bomb_pptx_is_refused_at_upload_and_never_decompressed(
    storage: FakeStorage, pptx_bytes: bytes, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Tetti ridotti per non scrivere centinaia di MB nel test: la logica è
    # la stessa (dimensioni dichiarate nello zip, controllate prima di leggere).
    monkeypatch.setattr(office, "MAX_XML_PART_BYTES", 64 * 1024)
    monkeypatch.setattr(office, "MAX_PACKAGE_BYTES", 2 * 1024 * 1024)
    bomb = _rewrite(pptx_bytes, {"ppt/slides/slide1.xml": b"<a/>" + b" " * (3 * 1024 * 1024)})
    assert len(bomb) < len(pptx_bytes) + 64 * 1024  # piccolo compresso, enorme decompresso
    with pytest.raises(ValidationAppError) as info:
        await file_service.save_document_from_bytes(bomb, subdir="courses/abc", mime_type=PPTX_MIME)
    assert info.value.code == "invalid_document_content"
    assert storage.files == {}
    # Anche se fosse già nello storage: testo e immagini rifiutano il pacchetto.
    source = tmp_path / "bomba.pptx"
    source.write_bytes(bomb)
    with pytest.raises(office.OfficeFormatError):
        office.pptx_text(str(source))
    with pytest.raises(office.OfficeFormatError):
        office.pptx_images(str(source))
    # Una sola slide oltre il tetto di parte (pacchetto sotto il tetto).
    big_slide = _rewrite(pptx_bytes, {"ppt/slides/slide1.xml": b"<a/>" + b" " * (200 * 1024)})
    source.write_bytes(big_slide)
    with pytest.raises(office.OfficeFormatError, match="parte troppo grande"):
        office.pptx_text(str(source))


def test_repeated_slide_ids_count_once(pptx_bytes: bytes, tmp_path: Path) -> None:
    presentation = _part(pptx_bytes, "ppt/presentation.xml")
    first = presentation.index("<p:sldId ")
    entry = presentation[first : presentation.index("/>", first) + 2]
    repeated = presentation.replace(entry, entry * 300, 1)
    source = tmp_path / "ripetuta.pptx"
    source.write_bytes(_rewrite(pptx_bytes, {"ppt/presentation.xml": repeated.encode("utf-8")}))
    original = tmp_path / "originale.pptx"
    original.write_bytes(pptx_bytes)
    assert office.pptx_slide_count(str(source)) == office.pptx_slide_count(str(original))
    assert office.pptx_text(str(source)) == office.pptx_text(str(original))


def test_absolute_relationship_targets_are_resolved(pptx_bytes: bytes, tmp_path: Path) -> None:
    rels = _part(pptx_bytes, "ppt/_rels/presentation.xml.rels")
    absolute = rels.replace('Target="slides/', 'Target="/ppt/slides/')
    assert absolute != rels
    source = tmp_path / "assoluti.pptx"
    source.write_bytes(
        _rewrite(pptx_bytes, {"ppt/_rels/presentation.xml.rels": absolute.encode("utf-8")})
    )
    original = tmp_path / "originale.pptx"
    original.write_bytes(pptx_bytes)
    assert office.pptx_slide_count(str(source)) == office.pptx_slide_count(str(original)) > 0
    assert office.pptx_text(str(source)) == office.pptx_text(str(original))
    assert len(office.pptx_images(str(source))) == len(office.pptx_images(str(original)))


def test_docx_captions_above_the_pictures(tmp_path: Path) -> None:
    from docx import Document
    from docx.shared import Cm
    from PIL import Image, ImageDraw

    def picture(color: str) -> io.BytesIO:
        image = Image.new("RGB", (400, 300), "white")
        ImageDraw.Draw(image).rectangle([40, 40, 360, 260], outline=color, width=8)
        buf = io.BytesIO()
        image.save(buf, format="PNG")
        buf.seek(0)
        return buf

    document = Document()
    document.add_paragraph("Figura 1. Schema a blocchi della catena di misura.")
    document.add_picture(picture("red"), width=Cm(8))
    document.add_paragraph("Figura 2. Banco di misura con lo shaker.")
    document.add_picture(picture("blue"), width=Cm(8))
    document.add_paragraph("Testo successivo.")
    path = tmp_path / "sopra.docx"
    document.save(str(path))
    images = office.docx_images(str(path))
    assert [i.caption for i in images] == [
        "Figura 1. Schema a blocchi della catena di misura.",
        "Figura 2. Banco di misura con lo shaker.",
    ]


async def test_slide_text_feeds_the_summary(tmp_path: Path, pptx_bytes: bytes) -> None:
    source = tmp_path / "lezione.pptx"
    source.write_bytes(pptx_bytes)
    text, _chars = await document_extraction_service.extract_text(source, PPTX_MIME)
    for needle in GROUND_TRUTH["pptx"]["text"]:
        assert needle in text


async def test_worker_extracts_pictures_per_slide(
    seeded_db: AsyncSession,
    storage: FakeStorage,
    pptx_bytes: bytes,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    patched = get_settings().model_copy(
        update={"figure_extraction_enabled": True, "figure_extraction_engine": "heuristic"}
    )
    monkeypatch.setattr(worker, "get_settings", lambda: patched)
    monkeypatch.setattr(worker, "mem_available_mb", lambda: None)
    vision = FakeVision()
    monkeypatch.setattr(worker, "describe_figure", vision)

    course_id, _org, _user = await build_course(seeded_db, modules=1, lessons_per_module=1)
    doc = build_course_document(course_id, filename="lezione_vibrometria.pptx")
    doc.mime_type = PPTX_MIME
    doc.file_path = f"/uploads/courses/{course_id}/lezione.pptx"
    doc.figures_status = "pending"
    seeded_db.add(doc)
    await seeded_db.commit()
    storage.files[remote_storage.uploads_key(doc.file_path)] = pptx_bytes

    claimed = await worker.claim_next(seeded_db)
    assert claimed is not None and claimed.id == doc.id
    await worker.process_document(seeded_db, claimed)
    assert claimed.figures_status == "ready", claimed.figures_error_code
    assert claimed.figures_pages_total == 3 and claimed.figures_count == 2
    assert claimed.bibliography == {
        "title": "Lezione di vibrometria",
        "authors": ["Docente di Prova"],
    }
    rows = list(
        (
            await seeded_db.execute(
                select(CourseDocumentFigure)
                .where(CourseDocumentFigure.document_id == doc.id)
                .order_by(CourseDocumentFigure.locator)
            )
        )
        .scalars()
        .all()
    )
    assert [(r.locator, r.page, r.status) for r in rows] == [
        ("s0002-f01", 2, "ready"),
        ("s0003-f01", 3, "ready"),
    ]
    line: Any = figure_attribution_line(rows[0], claimed, language="it")
    assert line == "Fonte: Docente di Prova, «Lezione di vibrometria», fig. 1, slide 2"


def test_entity_expansion_bombs_are_refused_before_parsing(
    pptx_bytes: bytes, tmp_path: Path
) -> None:
    """Fase D: una slide con una DTD e migliaia di riferimenti a un'entità
    interna (17 KB caricati → GB di RAM nel processo principale durante il
    riassunto) viene rifiutata prima di costruire l'albero; Office Open XML
    non usa mai DTD. Stessa guardia per `core.xml` dei metadati."""
    entity = "A" * 250
    body = "&a;" * 20_000
    bomb_slide = (
        f'<?xml version="1.0"?><!DOCTYPE p:sld [<!ENTITY a "{entity}">]>'
        '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f"<p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r><a:t>{body}</a:t></a:r></a:p>"
        "</p:txBody></p:sp></p:spTree></p:cSld></p:sld>"
    ).encode()
    source = zipfile.ZipFile(io.BytesIO(pptx_bytes))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for info in source.infolist():
            data = source.read(info)
            if info.filename.startswith("ppt/slides/slide") and info.filename.endswith(".xml"):
                data = bomb_slide
            zf.writestr(info, data)
    path = tmp_path / "bomba.pptx"
    path.write_bytes(out.getvalue())
    with pytest.raises(office.OfficeFormatError, match="DTD"):
        office.pptx_text(str(path))
    with pytest.raises(office.OfficeFormatError, match="DTD"):
        office.safe_fromstring(b'<!DOCTYPE x [<!ENTITY a "b">]><x>&a;</x>')
    assert office.safe_fromstring(b"<x><y>ok</y></x>").find("y").text == "ok"


def test_doc_upload_with_a_dtd_is_refused_before_docx2txt(tmp_path: Path) -> None:
    """Fase D (confutatore): un «.doc» che è uno zip con una DTD finiva a
    `docx2txt` senza guardie (1 KB → 43 MB di testo)."""
    path = tmp_path / "bomba.doc"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr(
            "word/document.xml",
            '<!DOCTYPE w [<!ENTITY a "' + "A" * 200 + '">]><w>' + "&a;" * 5000 + "</w>",
        )
    with pytest.raises(document_extraction_service.DocumentExtractionError, match="DTD"):
        document_extraction_service._extract_doc(path)
